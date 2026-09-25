"""Global Mangrove Watch v4.1.12 mangrove extent by country, 1985–2025 (polygon contract) →
public/data/gmw.geojson.

Item 3a of the 2026-09-25 wave (grill_wildeye_next_wave_2026-09-25 Q11). Source: Bunting, Hilarides,
Rosenqvist et al. 2026, "Global Mangrove Watch: Timeseries of Mangrove Extent" v4.1.12, Zenodo,
doi:10.5281/zenodo.21346457, CC BY 4.0 (record read 2026-09-25). The country file gives annual extent
in hectares (already corrected ×0.9775 for the accuracy assessment) with lower/upper 95% bounds, one
sheet each. It is downloaded once, its md5 checked against the Zenodo record, and cached.

Shapes: Natural Earth 50m admin-0 MAP UNITS (public domain), which split overseas territories such as
French Guiana from France; they match all but 4 GMW territories (0.006% of 2025 extent, 2026-09-25).
Those are listed in the output's `missing`, never dropped silently.
"""

from __future__ import annotations
import argparse
import datetime as dt
import hashlib
import json
import logging
import os
import time
import urllib.request
from pathlib import Path

from .atomic import write_atomic

log = logging.getLogger("gmw")
RECORD = "21346457"
DOI = "10.5281/zenodo.21346457"
VERSION = "4.1.12"
XLSX = "gmw_v4_timeseries_4112_gmw_country_stats_corr_area_formatted.xlsx"
XLSX_URL = f"https://zenodo.org/api/records/{RECORD}/files/{XLSX}/content"
XLSX_MD5 = "30800462dcd06a184df33aa66e0337a2"  # Zenodo checksum, 2026-09-25
NE_URL = "https://raw.githubusercontent.com/nvkelso/natural-earth-vector/master/geojson/ne_50m_admin_0_map_units.geojson"
UA = "wildeye/0.1 (gmw sync)"
LICENCE = "CC BY 4.0 (Global Mangrove Watch v4.1.12, Zenodo record licence)"
YEARS = list(range(1985, 2026))
SHEETS = ("Extent", "lower95th", "upper95th")


def _sheet_rows(ws, sheet: str) -> list[tuple]:
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    if (
        list(header[:2]) != ["iso", "cnty_name"]
        or [int(h) for h in header[2:] if h is not None] != YEARS
    ):
        raise ValueError(
            f"{sheet}: expected iso, cnty_name and years {YEARS[0]}–{YEARS[-1]}, got {list(header)[:5]}…"
        )
    return [r for r in rows[1:] if r and r[0]]


def read_stats(path: Path) -> dict[str, dict]:
    """ISO → {name, ha, lo, hi} (41 values each). Any structural surprise or an extent outside its bounds raises."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    missing = [s for s in SHEETS if s not in wb.sheetnames]
    if missing:
        raise ValueError(f"sheets {missing} not in {wb.sheetnames}")
    ext, lo, hi = (_sheet_rows(wb[s], s) for s in SHEETS)
    out = {}
    for sheet, rows in (("lower95th", lo), ("upper95th", hi)):
        if [r[0] for r in rows] != [r[0] for r in ext]:
            raise ValueError(f"{sheet}: countries are not in the Extent sheet's order")
    for e, l, h in zip(ext, lo, hi):
        iso, name = e[0], e[1]
        ha, lv, hv = (list(map(float, r[2 : 2 + len(YEARS)])) for r in (e, l, h))
        for y, a, b, c in zip(YEARS, lv, ha, hv):
            if not (a <= b + 1e-9 and b <= c + 1e-9):
                raise ValueError(
                    f"{iso} {y}: extent {b} outside its 95% bounds [{a}, {c}]"
                )
        out[iso] = {"name": name, "ha": ha, "lo": lv, "hi": hv}
    return out


def _rnd(c, nd=2):
    if isinstance(c[0], (int, float)):
        return [round(c[0], nd), round(c[1], nd)]
    return [_rnd(x, nd) for x in c]


def load_units(path: Path) -> dict[str, dict]:
    """ISO3 → MultiPolygon geometry (coordinates rounded to 0.01°), every map unit of that ISO merged."""
    out: dict[str, dict] = {}
    for f in json.loads(path.read_text())["features"]:
        p = f["properties"]
        iso = (
            p.get("ISO_A3")
            if p.get("ISO_A3") not in (None, "-99")
            else p.get("ADM0_A3")
        )
        if not iso or iso == "-99":
            continue
        g = f["geometry"]
        polys = g["coordinates"] if g["type"] == "MultiPolygon" else [g["coordinates"]]
        out.setdefault(iso, {"type": "MultiPolygon", "coordinates": []})[
            "coordinates"
        ].extend(_rnd(polys))
    return out


def build(stats: dict, units: dict) -> tuple[list[dict], list[dict]]:
    """Features for every country with mangrove in some year and a shape; `missing` names the rest."""
    feats, missing = [], []
    for iso in sorted(stats):
        s = stats[iso]
        if max(s["ha"]) <= 0:
            continue
        if iso not in units:
            missing.append(
                {"iso": iso, "name": s["name"], "ha_last": round(s["ha"][-1], 1)}
            )
            continue
        r = lambda v: [round(x, 1) for x in v]
        feats.append(
            {
                "type": "Feature",
                "geometry": units[iso],
                "properties": {
                    "iso": iso,
                    "name": s["name"],
                    "ha": r(s["ha"]),
                    "lo": r(s["lo"]),
                    "hi": r(s["hi"]),
                },
            }
        )
    if missing:
        log.warning(
            "%d GMW territories without a Natural Earth map unit: %s",
            len(missing),
            [m["iso"] for m in missing],
        )
    return feats, missing


def _fetch_bytes(url: str) -> bytes:
    return urllib.request.urlopen(
        urllib.request.Request(url, headers={"User-Agent": UA}), timeout=300
    ).read()


def _cached(path: Path, url: str, fetch_bytes) -> Path:
    if not path.exists():
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(fetch_bytes(url))
    return path


def main(argv=None, *, fetch_bytes=_fetch_bytes, md5=XLSX_MD5):
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=Path("public/data/gmw.geojson"))
    ap.add_argument(
        "--cache",
        type=Path,
        default=Path(
            os.environ.get("WILDEYE_CACHE", Path.home() / ".cache" / "wildeye")
        ),
    )
    a = ap.parse_args(argv)
    logging.basicConfig(
        level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s"
    )
    t0 = time.time()
    xlsx = a.cache / f"gmw_{VERSION}_{XLSX}"
    if not xlsx.exists():
        data = fetch_bytes(XLSX_URL)
        got = hashlib.md5(data).hexdigest()
        if got != md5:
            raise SystemExit(
                f"{XLSX_URL}: md5 {got} does not match the Zenodo record's {md5}"
            )
        xlsx.parent.mkdir(parents=True, exist_ok=True)
        xlsx.write_bytes(data)
    stats = read_stats(xlsx)
    units = load_units(
        _cached(a.cache / "ne_50m_admin_0_map_units.geojson", NE_URL, fetch_bytes)
    )
    feats, missing = build(stats, units)
    if not feats:
        raise SystemExit("no mangrove countries matched a shape")
    write_atomic(
        a.out,
        {
            "type": "FeatureCollection",
            "generated_at": dt.datetime.now(dt.UTC).strftime("%Y-%m-%dT%H:%M:%SZ"),
            "years": [YEARS[0], YEARS[-1]],
            "source": {
                "id": "gmw",
                "name": "Global Mangrove Watch",
                "version": VERSION,
                "doi": DOI,
                "licence": LICENCE,
                "url": f"https://doi.org/{DOI}",
                "shapes": "Natural Earth 50m admin-0 map units (public domain)",
                "citation": f"Bunting, P., Hilarides, L., Rosenqvist, A., et al. (2026). Global Mangrove Watch: Timeseries of Mangrove Extent, v{VERSION}. Zenodo. doi:{DOI}",
                "note": "Mangrove extent in hectares per country and year, corrected (×0.9775) for the accuracy assessment, with 95% bounds. Global F1 0.93.",
            },
            "missing": missing,
            "features": feats,
        },
    )
    log.info(
        "wrote %s: %d countries, %d without a shape (%.0f s)",
        a.out,
        len(feats),
        len(missing),
        time.time() - t0,
    )


if __name__ == "__main__":
    main()
